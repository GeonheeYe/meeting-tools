# ~/meeting_tools/context_loader.py
"""
참고 문서에서 키워드와 본문, 용어 메타데이터를 추출한다.
- key_terms: Whisper initial_prompt에 삽입할 쉼표 구분 핵심 용어
- doc_content: Claude 교정 단계에 전달할 문서 본문
- term_metadata: 용어 우선순위와 표기 정규화에 사용할 메타데이터
"""
import re
from collections import Counter
from pathlib import Path
from typing import Optional


def extract_text(path: Path) -> str:
    """파일 포맷에 맞게 텍스트 추출."""
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8")
    elif suffix == ".pdf":
        return _extract_pdf(path)
    elif suffix == ".docx":
        return _extract_docx(path)
    elif suffix in {".xlsx", ".xls"}:
        return _extract_xlsx(path)
    else:
        print(f"[context_loader] 지원하지 않는 포맷: {suffix}, 건너뜀")
        return ""


def _extract_pdf(path: Path) -> str:
    import pdfplumber
    pages = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
    return "\n".join(pages)


def _extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        rows.append(f"[시트: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
    wb.close()
    return "\n".join(rows)


def extract_agenda_items(text: str) -> list:
    """
    회의 안건 문서에서 최상위 항목 제목 목록을 추출한다.
    번호 패턴: '1.', '①', '(1)' 등으로 시작하는 들여쓰기 없는 줄.
    파싱 결과 0~1개이면 빈 리스트 반환 (안건 구조 없는 것으로 처리).
    """
    patterns = [
        r'^\d+\.\s+(.+)',       # 1. 항목
        r'^[①②③④⑤⑥⑦⑧⑨⑩]\s*(.+)',  # ① 항목
        r'^\(\d+\)\s+(.+)',     # (1) 항목
    ]
    items = []
    for line in text.splitlines():
        # 들여쓰기 있으면 하위 항목 → 건너뜀
        if line != line.lstrip():
            continue
        for pattern in patterns:
            m = re.match(pattern, line.strip())
            if m:
                items.append(m.group(1).strip())
                break

    if len(items) <= 1:
        return []
    return items


def extract_key_terms(text: str, max_terms: int = 30) -> str:
    """
    텍스트에서 핵심 용어 추출 (규칙 기반).
    - 영문 대문자로 시작하는 단어/약어 (BIRD, XQBot, GRPO 등)
    - 한글 명사구 (2자 이상 연속 한글)
    - 중복 제거, max_terms 개까지
    """
    en_terms = re.findall(r'\b[A-Z][a-zA-Z0-9\-]+\b|\b[A-Z]{2,}\b', text)
    ko_terms = re.findall(r'[가-힣]{2,}', text)

    all_terms = en_terms + ko_terms
    counted = Counter(all_terms)

    _STOPWORDS = {
        "있다", "없다", "하다", "된다", "이다", "합니다", "입니다",
        "있는", "없는", "하는", "되는", "이런", "그런", "저런",
        "The", "This", "That", "For", "With", "And", "Are",
        "활용합니다", "프레임워크", "벤치마크",
    }
    terms = [t for t, _ in counted.most_common(max_terms * 2)
             if t not in _STOPWORDS and len(t) >= 2]

    return ", ".join(terms[:max_terms])


def extract_term_metadata(text: str, max_terms: int = 30) -> dict:
    """문서 기반 용어 메타데이터를 추출한다."""
    raw_terms = []

    # 제품명/프로젝트명 후보: 대문자 약어, 하이픈 포함 용어, 영문+숫자 조합
    raw_terms.extend(re.findall(r"\b[A-Z]{2,}(?:-[A-Z0-9]+)+\b", text))
    raw_terms.extend(re.findall(r"\b[A-Z][A-Za-z0-9]+(?:-[A-Za-z0-9]+)?\b", text))

    # 한글 기반 기능명/과제명 후보
    raw_terms.extend(re.findall(r"[가-힣A-Za-z0-9·\-\[\]: ]{4,}(?:기능 개발|분석 기능|파이프라인|목표합의서|기획 및 설계|고도화)", text))

    cleaned_terms = []
    for term in raw_terms:
        cleaned = " ".join(term.split()).strip(" ,|")
        if len(cleaned) >= 2:
            cleaned_terms.append(cleaned)

    counted = Counter(cleaned_terms)
    canonical_terms = [term for term, _ in counted.most_common(max_terms)]

    alias_map = {}
    for term in canonical_terms:
        normalized = re.sub(r"[\s\-_:·\[\]\(\)]", "", term).lower()
        if normalized and normalized != term.lower():
            alias_map[normalized] = term
        if "-" in term:
            alias_map[term.replace("-", " ")] = term
            alias_map[term.replace("-", "")] = term

    for acronym in re.findall(r"\b[A-Z]{2,}\b", text):
        alias_map[acronym.title()] = acronym

    return {
        "canonical_terms": canonical_terms,
        "priority_terms": canonical_terms[:10],
        "alias_map": alias_map,
    }


def load(doc_paths: list) -> tuple:
    """
    문서 목록을 로드해서 (key_terms, doc_content, term_metadata) 반환.
    - key_terms: initial_prompt에 삽입
    - doc_content: Claude 교정용 전문
    """
    texts = []
    for p in doc_paths:
        path = Path(p).expanduser().resolve()
        if not path.exists():
            print(f"[context_loader] 파일 없음: {path}, 건너뜀")
            continue
        text = extract_text(path)
        if text:
            texts.append(f"=== {path.name} ===\n{text}")

    if not texts:
        return "", "", {"canonical_terms": [], "priority_terms": [], "alias_map": {}}

    doc_content = "\n\n".join(texts)
    key_terms = extract_key_terms(doc_content)
    term_metadata = extract_term_metadata(doc_content)
    print(f"[context_loader] 키워드 추출 완료: {key_terms[:80]}...")
    return key_terms, doc_content, term_metadata


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("사용법: python context_loader.py <파일1> [파일2] ...")
        sys.exit(1)
    terms, content, term_metadata = load(sys.argv[1:])
    print(f"\n=== 키워드 ===\n{terms}")
    print(f"\n=== 본문 (앞 500자) ===\n{content[:500]}")
    print(f"\n=== 우선 용어 ===\n{term_metadata['priority_terms']}")
